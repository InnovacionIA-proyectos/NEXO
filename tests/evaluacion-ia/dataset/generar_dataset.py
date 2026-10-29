"""Genera el dataset dorado de mapeo de Excel (S6-HU-32.5, #353): réplicas anonimizadas y la respuesta correcta de cada caso.

Uso (solo en el equipo que tiene los Excels originales):
    python generar_dataset.py            todos los casos
    python generar_dataset.py caso_05_…  solo los casos indicados

Lee de qué archivo y hoja sale cada caso en fuentes.local.json (no se versiona: los nombres de archivo y de hoja identifican al
cliente). Ver fuentes.local.example.json y el README.
"""
import datetime
import io
import json
import random
import re
import sys
import unicodedata
from pathlib import Path

import openpyxl

AQUI = Path(__file__).resolve().parent
sys.path.insert(0, str(AQUI))
from analizar_excel import analizar  # noqa: E402

RAIZ_REPO = AQUI.parents[2]
MUNICIPIOS = RAIZ_REPO / 'docs' / 'Database' / 'seeder' / 'municipios.csv'

# Tratamiento de cada columna en la réplica, por índice de encabezado (el que produce analizar-excel):
#   mantener      vocabulario que el modelo debe reconocer (vehículos, consecutivos, zonas): se copia tal cual
#   ciudad        se reemplaza por una ciudad colombiana del catálogo (misma ciudad -> mismo reemplazo dentro del caso)
#   ruta          cadena de ciudades separadas por guion: cada ciudad se reemplaza
#   departamento  el departamento de la ciudad reemplazada en la columna indicada
#   numero        (por defecto para celdas numéricas) número inventado de magnitud parecida
# Las notas largas y los textos de arriba del encabezado siempre se reemplazan por texto genérico.

C = 'campo'


def col(campo='ignorar', transformacion='ninguna', nota=None, despivotar=False, alternativas=None):
    return {'campo': campo, 'transformacion': transformacion, 'nota': nota, 'requiereDespivotar': despivotar,
            'alternativasAceptadas': alternativas or []}


IGN = col()
DIV = [{'campo': 'destino', 'transformacion': 'dividir_origen_destino'}]

CASOS = {
    'caso_01_cliente_a_importacion': {
        'cliente': 'A', 'formato': 'vehiculo_en_encabezado',
        'descripcion': 'Importación desde puerto: una fila por ruta; el volumen y la tarifa vienen en una columna por clase de viaje y vehículo.',
        'tratamiento': {0: 'mantener', 1: 'ciudad', 2: 'ciudad'},
        'respuesta': {
            0: col('referencia', nota='Consecutivo del cliente (Item-001...).'),
            1: col('origen'), 2: col('destino'),
            **{i: col(nota='Volumen de viajes para la clase/vehículo del encabezado: no hay campo para el volumen y el vehículo viene en el encabezado.',
                      despivotar=True) for i in range(3, 9)},
            **{i: col(nota='Tarifa que diligencia el transportador.') for i in range(9, 19)},
            **{i: col(nota='Tiempo de tránsito que diligencia el transportador.') for i in range(19, 24)},
            24: col(nota='Tarifa genset que diligencia el transportador.'),
        },
        'limitaciones': ['El vehículo y la clase de viaje vienen en el encabezado de cada columna de volumen (requiere despivotar).'],
    },
    'caso_02_cliente_a_exportacion': {
        'cliente': 'A', 'formato': 'vehiculo_en_encabezado',
        'descripcion': 'Exportación terrestre: una fila por ruta; volumen y tarifa en una columna por vehículo. Destinos con sufijo -DTAI.',
        'tratamiento': {0: 'mantener', 1: 'ciudad', 2: 'ciudad'},
        'respuesta': {
            0: col('referencia', nota='Consecutivo del cliente.'),
            1: col('origen'), 2: col('destino', nota='Algunos destinos traen el sufijo -DTAI (tránsito aduanero).'),
            **{i: col(nota='Volumen por vehículo: el vehículo viene en el encabezado.', despivotar=True) for i in range(3, 7)},
            **{i: col(nota='Tarifa que diligencia el transportador.') for i in range(7, 11)},
            11: col(nota='Tiempo de tránsito que diligencia el transportador.'),
            12: col(nota='Descripción de actividades que diligencia el transportador.'),
        },
        'limitaciones': ['El vehículo viene en el encabezado de cada columna de volumen (requiere despivotar).'],
    },
    'caso_03_cliente_a_nacional': {
        'cliente': 'A', 'formato': 'vehiculo_en_encabezado',
        'descripcion': 'Nacional: una fila por ruta con departamento del destino; volumen y tarifa por vehículo con su capacidad.',
        'tratamiento': {0: 'mantener', 1: 'ciudad', 2: 'ciudad', 3: 'departamento:2'},
        'respuesta': {
            0: col('referencia', nota='Consecutivo del cliente.'),
            1: col('origen'), 2: col('destino'),
            3: col(nota='Departamento del destino: no hay campo propio (el destino ya viene en su columna).'),
            **{i: col(nota='Volumen por vehículo: el vehículo (con su capacidad) viene en el encabezado.', despivotar=True) for i in range(4, 8)},
            8: col(nota='Volumen de paqueteo en kg: el tipo de servicio viene en el encabezado.', despivotar=True),
            **{i: col(nota='Tarifa que diligencia el transportador.') for i in range(9, 15)},
            **{i: col(nota='Dato que diligencia el transportador.') for i in range(15, 20)},
        },
        'limitaciones': ['El vehículo viene en el encabezado de cada columna de volumen (requiere despivotar).'],
    },
    'caso_04_cliente_b_escenarios': {
        'cliente': 'B', 'formato': 'vehiculo_en_encabezado',
        'descripcion': 'Hoja de trabajo con dos escenarios en la misma hoja (dos tablas, la segunda con su propio encabezado dentro de los datos).',
        'tratamiento': {0: 'ciudad', 1: 'mantener', 2: 'ciudad', 5: 'mantener'},
        'respuesta': {
            0: col('origen'),
            1: col(nota='Zona comercial del destino; algunas filas traen país y vehículo ("Ecuador Mula").'),
            2: col('destino'),
            3: col(nota='Viajes por mes: no hay campo para la frecuencia.'),
            4: col(nota='Flete de referencia del cliente; el vehículo (Sencillo) viene en el encabezado.', despivotar=True),
            5: col(nota='Observaciones sin encabezado.'),
        },
        'limitaciones': ['El vehículo viene en el encabezado de la columna de flete.',
                         'La hoja trae dos tablas: la segunda (ESCENARIO 2) repite su encabezado dentro de los datos.'],
    },
    'caso_05_cliente_c_region_1': {
        'cliente': 'C', 'formato': 'matriz',
        'descripcion': 'Formulario en matriz: tipo de servicio/vehículo en filas y puerto de origen en columnas; el destino es la región del formulario.',
        'tratamiento': {1: 'mantener', 2: 'mantener'},
        # El original es el formulario ya diligenciado: las tarifas se vacían, como llega una licitación.
        'vaciar': range(3, 11),
        'respuesta': {
            0: col(nota='Columna vacía.'),
            1: col('tipoVehiculo', nota='Tipo de servicio: vehículos, contenedores por rango de peso y servicios adicionales mezclados.'),
            2: col(nota='Repite la columna anterior (celda combinada).', alternativas=[{'campo': 'tipoVehiculo', 'transformacion': 'ninguna'}]),
            **{i: col(nota='Tarifa por puerto de origen: el origen viene en el encabezado.', despivotar=True) for i in range(3, 11)},
            11: IGN, 12: IGN,
            13: col(nota='Valor de referencia sin encabezado.'),
        },
        'limitaciones': ['El origen viene en el encabezado de cada columna (requiere despivotar).',
                         'El destino no está en ninguna columna: es la región del formulario (título).'],
    },
    'caso_06_cliente_c_region_2': {
        'cliente': 'C', 'formato': 'matriz',
        'descripcion': 'Mismo formulario en matriz para otra región, con un puerto de origen más.',
        'tratamiento': {1: 'mantener', 2: 'mantener'},
        'respuesta': {
            0: col(nota='Columna vacía.'),
            1: col('tipoVehiculo', nota='Tipo de servicio: vehículos, contenedores por rango de peso y servicios adicionales mezclados.'),
            2: col(nota='Repite la columna anterior (celda combinada).', alternativas=[{'campo': 'tipoVehiculo', 'transformacion': 'ninguna'}]),
            **{i: col(nota='Tarifa por puerto de origen: el origen viene en el encabezado.', despivotar=True) for i in range(3, 13)},
            13: IGN, 14: IGN,
            15: col(nota='Valor de referencia sin encabezado.'),
        },
        'limitaciones': ['El origen viene en el encabezado de cada columna (requiere despivotar).',
                         'El destino no está en ninguna columna: es la región del formulario (título).'],
    },
    'caso_07_cliente_d_impo_expo': {
        'cliente': 'D', 'formato': 'fila_por_linea',
        'descripcion': 'Formulario de oferta de importación y exportación: una fila por ruta y contenedor con rango de peso.',
        'tratamiento': {0: 'mantener', 1: 'ruta', 4: 'mantener', 5: 'mantener'},
        'respuesta': {
            0: col('referencia', nota='Identificador de la ruta del cliente.'),
            1: col('origen', 'dividir_origen_destino', alternativas=DIV, nota='«ORIGEN-DESTINO» en una sola celda.'),
            2: col(nota='Número de entregas: no hay campo.'),
            3: col(nota='Kilómetros de referencia: no hay campo.'),
            4: col('claseContenedor', nota='Contenedor de 20 o 40 pies con su rango de peso (C20 de 0 a 8 Ton).',
                   alternativas=[{'campo': 'tipoVehiculo', 'transformacion': 'ninguna'}]),
            5: col(nota='Configuración SICE-TAC.', alternativas=[{'campo': 'tipoVehiculo', 'transformacion': 'ninguna'}]),
            6: col(nota='Tiempo logístico de referencia de SICE-TAC.'),
            **{i: col(nota='Costo o tarifa que diligencia el transportador.') for i in range(7, 14)},
        },
        'limitaciones': ['El tipo de contenedor y el rango de peso vienen juntos en la misma celda.'],
    },
}

NOTA = re.compile(r'^\s*nota\s*\d|^\s*observaci', re.I)
_ACENTOS = {'a': '[aá]', 'e': '[eé]', 'i': '[ií]', 'o': '[oó]', 'u': '[uúü]', 'n': '[nñ]'}


def sin_tildes(s):
    return ''.join(c for c in unicodedata.normalize('NFD', s) if unicodedata.category(c) != 'Mn')


def cargar_municipios():
    import csv
    with io.open(MUNICIPIOS, encoding='utf-8-sig') as f:
        filas = [r for r in csv.DictReader(f) if r['Pais'] == 'CO']
    return [(r['Municipio'], r['Departamento']) for r in filas]


class Anonimizador:
    def __init__(self, semilla, municipios):
        self.rnd = random.Random(semilla)
        self.municipios = municipios
        self.ciudades = {}
        self.textos = {}
        self.notas = {}

    def ciudad(self, valor):
        v = str(valor).strip()
        pref = 'AEROPUERTO ' if sin_tildes(v).upper().startswith('AEROPUERTO ') else ''
        base = v[len(pref):]
        m = re.match(r'^(.*?)(-DTAI)?$', base, re.I)
        nombre, suf = m.group(1).strip(), m.group(2) or ''
        clave = sin_tildes(nombre).upper()
        if clave not in self.ciudades:
            usadas = {c for c, _ in self.ciudades.values()}
            libres = [m for m in self.municipios if m[0] not in usadas] or self.municipios
            self.ciudades[clave] = self.rnd.choice(libres)
        nuevo = self.ciudades[clave][0]
        nuevo = nuevo if nombre.isupper() else nuevo.title()
        return (pref.title() if pref and not v.isupper() else pref) + nuevo + suf

    def departamento_de(self, valor_original):
        clave = sin_tildes(re.sub(r'(?i)-DTAI$', '', str(valor_original).strip())).upper()
        return self.ciudades.get(clave, ('', ''))[1]

    def ruta(self, valor):
        partes = re.split(r'(\s*-\s*)', str(valor))
        return ''.join(p if i % 2 else self.ciudad(p) for i, p in enumerate(partes) if p != '')

    def numero(self, v):
        if isinstance(v, bool) or v == 0:
            return v
        x = v * self.rnd.uniform(0.6, 1.4)
        if isinstance(v, int):
            x = max(1, round(x, -3)) if abs(v) >= 10000 else max(1, round(x))
            if x == v:
                x = v + (1000 if abs(v) >= 10000 else 1)
            return x
        x = round(x, 2)
        return x if x != v else round(v + 0.5, 2)

    def ciudades_en_texto(self, texto):
        """Reemplaza, dentro de un texto libre, las ciudades que ya se reemplazaron en las columnas de ciudad."""
        for clave, (nuevo, _) in self.ciudades.items():
            patron = ''.join(_ACENTOS.get(ch, re.escape(ch)) for ch in clave.lower())
            texto = re.sub(r'(?i)\b%s\b' % patron,
                           lambda m: nuevo if m.group(0).isupper() else nuevo.title(), texto)
        return texto

    def nota(self, valor):
        if valor not in self.notas:
            self.notas[valor] = 'Nota %d: texto de ejemplo para la prueba.' % (len(self.notas) + 1)
        return self.notas[valor]

    def texto(self, valor):
        if valor not in self.textos:
            self.textos[valor] = 'Texto del formato %d' % (len(self.textos) + 1)
        return self.textos[valor]


def es_numero(v):
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def generar_caso(caso_id, spec, fuente, municipios):
    wb = openpyxl.load_workbook(RAIZ_REPO / fuente['archivo'], data_only=True)
    ws = wb[fuente['hoja']]
    original = analizar(ws)
    fe, fs = original['filaEncabezado'], original['filaSubEncabezado'] or original['filaEncabezado']
    pc = openpyxl.utils.column_index_from_string(original['primeraColumna'])

    # Celdas que forman el encabezado: las de sus filas y las combinadas que lo alcanzan desde arriba.
    protegidas = set()
    for rango in ws.merged_cells.ranges:
        if rango.min_row <= fs and rango.max_row >= fe:
            protegidas.add((rango.min_row, rango.min_col))
    for r in range(fe, fs + 1):
        for c in range(1, ws.max_column + 1):
            protegidas.add((r, c))

    an = Anonimizador(caso_id, municipios)
    trat = spec['tratamiento']
    originales = {}
    for fila in ws.iter_rows():
        for celda in fila:
            v = celda.value
            if v is None or (celda.row, celda.column) in protegidas or isinstance(celda, openpyxl.cell.cell.MergedCell):
                continue
            if celda.row < fe:
                celda.value = an.texto(v) if isinstance(v, str) and v.strip() else v
                continue
            indice = celda.column - pc
            t = trat.get(indice, '')
            originales[(celda.row, indice)] = v
            if indice in spec.get('vaciar', ()):
                celda.value = None
                continue
            if isinstance(v, str) and NOTA.match(v):
                celda.value = an.nota(v) if len(v) > 20 else v
            elif isinstance(v, str) and len(v) > 60:
                celda.value = an.nota(v)
            elif t == 'mantener':
                continue
            elif t == 'ciudad' and isinstance(v, str) and v.strip() and not re.match(r'(?i)^(escenario|origen|destino)\b', v.strip()):
                celda.value = an.ciudad(v)
            elif t == 'ruta' and isinstance(v, str):
                celda.value = an.ruta(v)
            elif es_numero(v):
                celda.value = an.numero(v)
            elif isinstance(v, (datetime.date, datetime.datetime)):
                celda.value = v
    for (r, indice), v in originales.items():
        t = trat.get(indice, '')
        celda = ws.cell(r, pc + indice)
        if t not in ('ciudad', 'ruta') and isinstance(celda.value, str) and not isinstance(celda, openpyxl.cell.cell.MergedCell):
            celda.value = an.ciudades_en_texto(celda.value)
        if t.startswith('departamento:') and isinstance(v, str) and v.strip():
            ref = int(t.split(':')[1])
            nuevo = an.departamento_de(originales.get((r, ref), ''))
            if nuevo:
                ws.cell(r, pc + indice).value = nuevo

    for otra in [h for h in wb.sheetnames if h != ws.title]:
        del wb[otra]
    ws.title = 'Datos'
    ws.sheet_state = 'visible'
    ws.data_validations.dataValidation = []
    for fila in ws.iter_rows():
        for celda in fila:
            if not isinstance(celda, openpyxl.cell.cell.MergedCell):
                celda.comment = None
                celda.hyperlink = None
    for nombre in list(wb.defined_names):
        del wb.defined_names[nombre]
    wb.properties.creator = 'NEXO'
    wb.properties.lastModifiedBy = 'NEXO'
    wb.properties.title = caso_id
    wb.properties.company = None
    wb.properties.description = None
    wb.properties.subject = None
    wb.properties.keywords = None

    salida = AQUI / 'casos' / (caso_id + '.xlsx')
    wb.save(salida)

    replica = analizar(openpyxl.load_workbook(salida, data_only=True)['Datos'])
    if replica['encabezados'] != original['encabezados'] or replica['filaEncabezado'] != fe:
        raise SystemExit('%s: la réplica cambió los encabezados que detecta analizar-excel' % caso_id)
    return replica


def respuesta(caso_id, spec, replica):
    enc = replica['encabezados']
    faltan = [i for i in range(len(enc)) if i not in spec['respuesta']]
    sobran = [i for i in spec['respuesta'] if i >= len(enc)]
    if faltan or sobran:
        raise SystemExit('%s: la respuesta no cubre exactamente las columnas (faltan %s, sobran %s)' % (caso_id, faltan, sobran))
    columnas = []
    for i, texto in enumerate(enc):
        r = dict(spec['respuesta'][i])
        columnas.append({'indice': i, 'columna': texto, 'campo': r['campo'], 'transformacion': r['transformacion'],
                         'requiereDespivotar': r['requiereDespivotar'], 'alternativasAceptadas': r['alternativasAceptadas'],
                         'nota': r['nota']})
    return {
        'caso': caso_id,
        'archivo': 'casos/%s.xlsx' % caso_id,
        'hoja': 'Datos',
        'cliente': spec['cliente'],
        'formato': spec['formato'],
        'descripcion': spec['descripcion'],
        'filaEncabezado': replica['filaEncabezado'],
        'filaSubEncabezado': replica['filaSubEncabezado'],
        'totalFilas': len(replica['filas']),
        'limitaciones': spec['limitaciones'],
        'columnas': columnas,
    }


def main():
    fuentes = json.loads((AQUI / 'fuentes.local.json').read_text(encoding='utf-8'))
    municipios = cargar_municipios()
    pedidos = sys.argv[1:] or list(CASOS)
    desconocidos = [c for c in pedidos if c not in CASOS]
    if desconocidos:
        raise SystemExit('Casos desconocidos: %s' % ', '.join(desconocidos))
    for caso_id in pedidos:
        spec = CASOS[caso_id]
        replica = generar_caso(caso_id, spec, fuentes[caso_id], municipios)
        doc = respuesta(caso_id, spec, replica)
        (AQUI / 'respuestas' / (caso_id + '.json')).write_text(
            json.dumps(doc, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
        utiles = sum(1 for c in doc['columnas'] if c['campo'] != 'ignorar')
        print('%-34s %3d columnas (%d con campo) %4d filas  %s' % (caso_id, len(doc['columnas']), utiles, doc['totalFilas'], doc['formato']))


if __name__ == '__main__':
    main()
